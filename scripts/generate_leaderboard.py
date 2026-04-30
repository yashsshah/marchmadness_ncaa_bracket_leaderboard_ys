"""
Generate the leaderboard JSON for all 63 participants.
Extracts real picks from the master xlsx and scores against verified actual results.
Scoring rules from NCAA_2026_Rules.doc.
"""
import json
import openpyxl
import re

MASTER_FILE = "files/2026_Bauman-Rehmer_NCAA.xlsx"
LEADERBOARD_FILE = "data/leaderboard.json"
COMPARISON_FILE = "data/bracket_comparisons.json"

# ===========================================================
# SCORING RULES (from NCAA_2026_Rules.doc)
# ===========================================================
# 1st round (R64) = 1 pt each  (32 games, max 32)
# 2nd round (R32) = 2 pt each  (16 games, max 32)
# 3rd round (S16) = 4 pt each  (8 games, max 32)
# 4th round (E8)  = 8 pt each  (4 games, max 32)
# Semi-Final (F4) = 10 pt each (2 games, max 20)
# Champion        = 12 pt      (1 game,  max 12)
# TOTAL MAX = 160

ROUND_POINTS = {"R64": 1, "R32": 2, "S16": 4, "E8": 8, "F4": 10, "NCG": 12}
MAX_SCORE = 160
REGIONS = ["East", "West", "South", "Midwest"]
COMPLETED_ROUNDS = ["R64", "R32", "S16", "E8"]

# ===========================================================
# ACTUAL RESULTS by bracket cell position
# ===========================================================
ACTUALS_BY_POS = {}

# --- EAST R64 (col B=2, rows 5,7,9,11,13,15,17,19) ---
ACTUALS_BY_POS[(5, 2)]  = "Duke"
ACTUALS_BY_POS[(7, 2)]  = "TCU"
ACTUALS_BY_POS[(9, 2)]  = "St. John's"
ACTUALS_BY_POS[(11, 2)] = "Kansas"
ACTUALS_BY_POS[(13, 2)] = "Louisville"
ACTUALS_BY_POS[(15, 2)] = "Michigan State"
ACTUALS_BY_POS[(17, 2)] = "UCLA"
ACTUALS_BY_POS[(19, 2)] = "UConn"

# --- EAST R32 (col C=3) ---
ACTUALS_BY_POS[(5, 3)]  = "Duke"
ACTUALS_BY_POS[(9, 3)]  = "St. John's"
ACTUALS_BY_POS[(13, 3)] = "Michigan State"
ACTUALS_BY_POS[(17, 3)] = "UConn"

# --- EAST S16 (col D=4) ---
ACTUALS_BY_POS[(5, 4)]  = "Duke"
ACTUALS_BY_POS[(13, 4)] = "UConn"

# --- EAST E8 (col E=5) ---
ACTUALS_BY_POS[(5, 5)]  = "UConn"

# --- WEST R64 (col L=12, rows 5,7,9,11,13,15,17,19) ---
ACTUALS_BY_POS[(5, 12)]  = "Arizona"
ACTUALS_BY_POS[(7, 12)]  = "Utah State"
ACTUALS_BY_POS[(9, 12)]  = "High Point"
ACTUALS_BY_POS[(11, 12)] = "Arkansas"
ACTUALS_BY_POS[(13, 12)] = "Texas"
ACTUALS_BY_POS[(15, 12)] = "Gonzaga"
ACTUALS_BY_POS[(17, 12)] = "Miami (FL)"
ACTUALS_BY_POS[(19, 12)] = "Purdue"

# --- WEST R32 (col K=11) ---
ACTUALS_BY_POS[(5, 11)]  = "Arizona"
ACTUALS_BY_POS[(9, 11)]  = "Arkansas"
ACTUALS_BY_POS[(13, 11)] = "Texas"
ACTUALS_BY_POS[(17, 11)] = "Purdue"

# --- WEST S16 (col J=10) ---
ACTUALS_BY_POS[(5, 10)]  = "Arizona"
ACTUALS_BY_POS[(13, 10)] = "Purdue"

# --- WEST E8 (col I=9) ---
ACTUALS_BY_POS[(5, 9)]   = "Arizona"

# --- SOUTH R64 (col B=2, rows 24,26,28,30,32,34,36,38) ---
ACTUALS_BY_POS[(24, 2)] = "Florida"
ACTUALS_BY_POS[(26, 2)] = "Iowa"
ACTUALS_BY_POS[(28, 2)] = "Vanderbilt"
ACTUALS_BY_POS[(30, 2)] = "Nebraska"
ACTUALS_BY_POS[(32, 2)] = "VCU"
ACTUALS_BY_POS[(34, 2)] = "Illinois"
ACTUALS_BY_POS[(36, 2)] = "Texas A&M"
ACTUALS_BY_POS[(38, 2)] = "Houston"

# --- SOUTH R32 (col C=3) ---
ACTUALS_BY_POS[(24, 3)] = "Iowa"
ACTUALS_BY_POS[(28, 3)] = "Nebraska"
ACTUALS_BY_POS[(32, 3)] = "Illinois"
ACTUALS_BY_POS[(36, 3)] = "Houston"

# --- SOUTH S16 (col D=4) ---
ACTUALS_BY_POS[(24, 4)] = "Iowa"
ACTUALS_BY_POS[(32, 4)] = "Illinois"

# --- SOUTH E8 (col E=5) ---
ACTUALS_BY_POS[(24, 5)] = "Illinois"

# --- MIDWEST R64 (col L=12, rows 24,26,28,30,32,34,36,38) ---
ACTUALS_BY_POS[(24, 12)] = "Michigan"
ACTUALS_BY_POS[(26, 12)] = "Saint Louis"
ACTUALS_BY_POS[(28, 12)] = "Texas Tech"
ACTUALS_BY_POS[(30, 12)] = "Alabama"
ACTUALS_BY_POS[(32, 12)] = "Tennessee"
ACTUALS_BY_POS[(34, 12)] = "Virginia"
ACTUALS_BY_POS[(36, 12)] = "Kentucky"
ACTUALS_BY_POS[(38, 12)] = "Iowa State"

# --- MIDWEST R32 (col K=11) ---
ACTUALS_BY_POS[(24, 11)] = "Michigan"
ACTUALS_BY_POS[(28, 11)] = "Alabama"
ACTUALS_BY_POS[(32, 11)] = "Tennessee"
ACTUALS_BY_POS[(36, 11)] = "Iowa State"

# --- MIDWEST S16 (col J=10) ---
ACTUALS_BY_POS[(24, 10)] = "Michigan"
ACTUALS_BY_POS[(32, 10)] = "Tennessee"

# --- MIDWEST E8 (col I=9) ---
ACTUALS_BY_POS[(24, 9)] = "Michigan"

FINAL_FOUR_TEAMS = ["UConn", "Arizona", "Illinois", "Michigan"]
FINALISTS = ["UConn", "Michigan"]
ACTUALS_BY_POS[(5, 6)] = "UConn"
ACTUALS_BY_POS[(5, 8)] = "Michigan"

# ===========================================================
# TEAM NAME NORMALIZATION
# ===========================================================
def normalize(name):
    if name is None:
        return ""
    n = str(name).strip()
    if not n:
        return ""
    # Remove seed prefix like "(1) " or "(16) "
    n = re.sub(r'^\(\d+\)\s*', '', n)
    # Lowercase
    n = n.lower()
    # Remove punctuation
    n = n.replace(".", "").replace("'", "").replace("\u2019", "").replace(",", "")
    n = n.replace("<", "").replace(">", "")  # Fix typos like "M<ichigan"
    n = n.strip()

    # PRIORITY-ORDERED aliases (longer/more specific patterns first)
    # Multi-word patterns checked before single-word
    ALIASES = [
        # --- Compound names (must match before single-word) ---
        # Michigan State variants
        ("michitgan st", "michiganstate"), ("mich st", "michiganstate"),
        ("michigan st", "michiganstate"), ("michigan state", "michiganstate"),
        ("mi st", "michiganstate"), ("msu", "michiganstate"),
        # Ohio State variants
        ("rjio st", "ohiostate"), ("ohio st", "ohiostate"),
        ("ohio state", "ohiostate"), ("osu", "ohiostate"),
        # Iowa State variants
        ("iowa st", "iowastate"), ("iowa state", "iowastate"),
        ("ia st", "iowastate"), ("isu", "iowastate"),
        # Utah State variants
        ("utah st", "utahstate"), ("utah state", "utahstate"), ("usu", "utahstate"),
        # Tennessee State
        ("tenn st", "tennesseestate"), ("tennessee st", "tennesseestate"),
        ("tennessee state", "tennesseestate"),
        # North Dakota State
        ("nd state", "northdakotastate"), ("north dakota st", "northdakotastate"),
        ("north dakota state", "northdakotastate"),
        # Kennesaw State
        ("kennesaw st", "kennesawstate"), ("kennesaw state", "kennesawstate"),
        # Wright State
        ("wright st", "wrightstate"), ("wright state", "wrightstate"),
        # St. John's variants
        ("st johns", "stjohns"), ("st john", "stjohns"), ("saint johns", "stjohns"),
        ("st, johns", "stjohns"),
        # St. Mary's variants
        ("st marys", "saintmarys"), ("saint marys", "saintmarys"),
        ("st mary", "saintmarys"),
        # St. Louis
        ("st louis", "saintlouis"), ("saint louis", "saintlouis"),
        # Miami variants (specific first)
        ("miami of florida", "miamiflorida"), ("miami (florida)", "miamiflorida"),
        ("miami (fl)", "miamiflorida"), ("miami fl", "miamiflorida"),
        ("miami  oh", "miamiohio"), ("miami (oh)", "miamiohio"),
        ("miami (ohio)", "miamiohio"), ("miami oh", "miamiohio"),
        # Note: bare "miami" handled separately below
        # Texas variants (compound)
        ("texas a&m", "texasam"), ("texas am", "texasam"),
        ("tx a&m", "texasam"), ("tx am", "texasam"),
        ("t a&m", "texasam"), ("t am", "texasam"),
        ("tamu", "texasam"),
        ("texas tech", "texastech"), ("tex tech", "texastech"),
        ("tt", "texastech"), ("ttu", "texastech"),
        ("tx/ncst", "texas"),
        # North Carolina variants
        ("north caroina", "northcarolina"), ("north carolina", "northcarolina"),
        ("n carolina", "northcarolina"), ("ncarolina", "northcarolina"),
        ("nc", "northcarolina"),
        # Northern Iowa
        ("northern iowa", "northerniowa"), ("n iowa", "northerniowa"),
        ("uni", "northerniowa"),
        # South Florida
        ("south florida", "southflorida"), ("s florida", "southflorida"),
        ("usf", "southflorida"),
        # UConn variants
        ("u conn", "uconn"), ("uconn", "uconn"), ("conn", "uconn"),
        # High Point
        ("high point", "highpoint"), ("high pt", "highpoint"),
        # Michigan (bare) - AFTER Michigan State
        ("michigan", "michigan"), ("mich", "michigan"), ("mi", "michigan"),
        ("michitgan", "michigan"), ("m ichigan", "michigan"), ("mch", "michigan"),
        # Long Island
        ("long island", "longisland"), ("liu", "longisland"),
        # Cal Baptist
        ("cal baptist", "californiabaptist"), ("california baptist", "californiabaptist"),
        # Prairie View
        ("prairie view a&m", "prairieview"), ("pv a&m", "prairieview"),
        ("prairie view am", "prairieview"), ("pv a&m/lehigh", "prairieview"),
        # UMBC/Howard
        ("umbc / howard", "howard"), ("umbc/howard", "howard"),
        # TX / NC St (First Four: Texas beat NC State)
        ("tx / nc st", "texas"), ("tx / nc st", "texas"),
        # Miami OH / SMU (First Four)
        ("miami (oh) / smu", "miamiohio"), ("miami oh / smu", "miamiohio"),
        # Mc Neese
        ("mc neese", "mcneese"),
        # Queens
        ("queens (nc)", "queens"), ("queens nc", "queens"),
        # --- Single-word abbreviations ---
        ("unc", "northcarolina"),
        ("duke", "duke"),
        ("tcu", "tcu"),
        ("ku", "kansas"), ("kanas", "kansas"), ("kansas", "kansas"),
        ("ul", "louisville"), ("loisville", "louisville"), ("louisville", "louisville"),
        ("ucla", "ucla"),
        ("arizona", "arizona"), ("zona", "arizona"), ("az", "arizona"),
        ("arkansas", "arkansas"), ("ark", "arkansas"),
        ("gonzaga", "gonzaga"), ("zags", "gonzaga"),
        ("purdue", "purdue"), ("purdu", "purdue"),
        ("florida", "florida"), ("uf", "florida"), ("fl", "florida"),
        ("iowa", "iowa"),
        ("vanderbilt", "vanderbilt"), ("vandy", "vanderbilt"),
        ("nebraska", "nebraska"), ("ne", "nebraska"), ("nebrasjka", "nebraska"),
        ("vcu", "vcu"),
        ("illinois", "illinois"), ("il", "illinois"), ("illinoia", "illinois"),
        ("houston", "houston"), ("uh", "houston"),
        ("alabama", "alabama"), ("bama", "alabama"), ("al", "alabama"),
        ("tennessee", "tennessee"), ("tn", "tennessee"), ("ut", "tennessee"),
        ("virginia", "virginia"), ("uva", "virginia"), ("va", "virginia"),
        ("viginia", "virginia"), ("virgina", "virginia"),
        ("kentucky", "kentucky"), ("uk", "kentucky"), ("ky", "kentucky"),
        ("krntucky", "kentucky"),
        ("byu", "byu"),
        ("villanova", "villanova"), ("nova", "villanova"),
        ("wisconsin", "wisconsin"), ("uw", "wisconsin"), ("wi", "wisconsin"),
        ("wiscy", "wisconsin"),
        ("georgia", "georgia"), ("uga", "georgia"),
        ("missouri", "missouri"),
        ("clemson", "clemson"),
        ("penn", "penn"),
        ("troy", "troy"),
        ("miami", "miamiflorida"),  # Default bare "miami" to FL (the 7-seed)
        ("hawaii", "hawaii"),
        ("hofstra", "hofstra"),
        ("idaho", "idaho"),
        ("ucf", "ucf"),
        ("texas", "texas"), ("tx", "texas"),
        ("sienna", "siena"), ("siena", "siena"),
        ("mcneese", "mcneese"),
        ("akron", "akron"),
        ("howard", "howard"),
        ("queens", "queens"),
        ("furman", "furman"),
        ("santa clara", "santaclara"),
        ("highpoint", "highpoint"),
    ]

    for pattern, canonical in ALIASES:
        if n == pattern:
            return canonical

    # Fallback: strip remaining punctuation and spaces
    n = n.replace("(", "").replace(")", "").replace("&", "").replace(" ", "")
    return n


def names_match(pick, actual):
    return normalize(pick) == normalize(actual)


# ===========================================================
# BRACKET CELL POSITIONS PER REGION
# ===========================================================
def get_pick_positions(region):
    positions = []
    if region == "East":
        base, side = 5, "left"
    elif region == "South":
        base, side = 24, "left"
    elif region == "West":
        base, side = 5, "right"
    elif region == "Midwest":
        base, side = 24, "right"
    else:
        return positions

    if side == "left":
        for i in range(8):
            positions.append((base + i * 2, 2, "R64"))
        for i in range(4):
            positions.append((base + i * 4, 3, "R32"))
        for i in range(2):
            positions.append((base + i * 8, 4, "S16"))
        positions.append((base, 5, "E8"))
    else:
        for i in range(8):
            positions.append((base + i * 2, 12, "R64"))
        for i in range(4):
            positions.append((base + i * 4, 11, "R32"))
        for i in range(2):
            positions.append((base + i * 8, 10, "S16"))
        positions.append((base, 9, "E8"))
    return positions


F4_POSITIONS = [(5, 6, "F4"), (5, 8, "F4")]
NCG_POSITION = (5, 7, "NCG")

# Reverse lookup: canonical form → display name
DISPLAY_NAMES = {
    "duke": "Duke", "tcu": "TCU", "stjohns": "St. John's", "kansas": "Kansas",
    "louisville": "Louisville", "michiganstate": "Michigan State", "ucla": "UCLA",
    "uconn": "UConn", "arizona": "Arizona", "utahstate": "Utah State",
    "highpoint": "High Point", "arkansas": "Arkansas", "texas": "Texas",
    "gonzaga": "Gonzaga", "miamiflorida": "Miami (FL)", "purdue": "Purdue",
    "florida": "Florida", "iowa": "Iowa", "vanderbilt": "Vanderbilt",
    "nebraska": "Nebraska", "vcu": "VCU", "illinois": "Illinois",
    "texasam": "Texas A&M", "houston": "Houston", "michigan": "Michigan",
    "saintlouis": "Saint Louis", "texastech": "Texas Tech", "alabama": "Alabama",
    "tennessee": "Tennessee", "virginia": "Virginia", "kentucky": "Kentucky",
    "iowastate": "Iowa State", "byu": "BYU", "villanova": "Villanova",
    "wisconsin": "Wisconsin", "georgia": "Georgia", "missouri": "Missouri",
    "northcarolina": "North Carolina", "ohiostate": "Ohio State",
    "southflorida": "South Florida", "northerniowa": "Northern Iowa",
    "saintmarys": "Saint Mary's", "santaclara": "Santa Clara",
    "ucf": "UCF", "siena": "Siena", "furman": "Furman",
    "californiabaptist": "California Baptist", "northdakotastate": "North Dakota State",
    "hawaii": "Hawaii", "queens": "Queens", "prairieview": "Prairie View A&M",
    "mcneese": "McNeese", "hofstra": "Hofstra", "akron": "Akron",
    "wrightstate": "Wright State", "tennesseestate": "Tennessee State",
    "longisland": "Long Island", "miamiohio": "Miami (OH)", "howard": "Howard",
    "penn": "Penn", "troy": "Troy",
}


def clean_team_name(raw):
    """Strip seed prefix and return proper display name."""
    if not raw:
        return ""
    canonical = normalize(raw)
    return DISPLAY_NAMES.get(canonical, raw.strip())


# ===========================================================
# EXTRACT PICKS FROM A WORKSHEET
# ===========================================================
def extract_picks(ws):
    picks = {}
    for region in ["East", "West", "South", "Midwest"]:
        picks[region] = {}
        for pos in get_pick_positions(region):
            row, col, rd = pos
            val = ws.cell(row, col).value
            if rd not in picks[region]:
                picks[region][rd] = []
            picks[region][rd].append(val if val else "")

    picks["F4"] = []
    for row, col, _ in F4_POSITIONS:
        val = ws.cell(row, col).value
        picks["F4"].append(val if val else "")

    row, col, _ = NCG_POSITION
    picks["NCG"] = ws.cell(row, col).value or ""

    tb = ws.cell(44, 2).value
    if tb is not None:
        # Handle formats like "127(69-58)" or "127" or just a number
        tb_str = str(tb).strip()
        match = re.match(r'(\d+)', tb_str)
        picks["tiebreaker"] = int(match.group(1)) if match else 0
    else:
        picks["tiebreaker"] = 0

    return picks


# ===========================================================
# SCORE A SET OF PICKS
# ===========================================================
def score_picks(picks):
    round_correct = {"R64": 0, "R32": 0, "S16": 0, "E8": 0, "F4": 0, "NCG": 0}
    region_scores = {"East": 0, "West": 0, "South": 0, "Midwest": 0}

    for region in ["East", "West", "South", "Midwest"]:
        for pos in get_pick_positions(region):
            row, col, rd = pos
            actual = ACTUALS_BY_POS.get((row, col))
            if actual is None:
                continue
            positions_for_rd = [p for p in get_pick_positions(region) if p[2] == rd]
            idx = positions_for_rd.index(pos)
            pick_list = picks[region].get(rd, [])
            if idx < len(pick_list) and names_match(pick_list[idx], actual):
                round_correct[rd] += 1
                region_scores[region] += ROUND_POINTS[rd]

    champion = picks.get("NCG", "")
    f4_picks = picks.get("F4", ["", ""])

    # Clean champion name (strip seed prefix, normalize display)
    champion = clean_team_name(champion)
    f4_picks = [clean_team_name(p) for p in f4_picks]

    for idx, (row, col, rd) in enumerate(F4_POSITIONS):
        actual = ACTUALS_BY_POS.get((row, col))
        if actual is None:
            continue
        if idx < len(f4_picks) and names_match(f4_picks[idx], actual):
            round_correct[rd] += 1

    ncg_actual = ACTUALS_BY_POS.get((NCG_POSITION[0], NCG_POSITION[1]))
    if ncg_actual and champion and names_match(champion, ncg_actual):
        round_correct["NCG"] += 1

    round_scores = {rd: round_correct[rd] * ROUND_POINTS[rd] for rd in round_correct}
    total = sum(round_scores.values())

    champ_alive = bool(champion and normalize(champion) in [normalize(t) for t in FINALISTS])
    f4_alive = [p for p in f4_picks if p and normalize(p) in [normalize(t) for t in FINALISTS]]
    max_future = 0 if ncg_actual else (ROUND_POINTS["NCG"] if champ_alive else 0)
    max_possible = total + max_future

    return {
        "roundCorrect": [round_correct["R64"], round_correct["R32"],
                         round_correct["S16"], round_correct["E8"],
                         round_correct["F4"], round_correct["NCG"]],
        "roundScores": [round_scores["R64"], round_scores["R32"],
                        round_scores["S16"], round_scores["E8"],
                        round_scores["F4"], round_scores["NCG"]],
        "regionScores": region_scores,
        "totalPoints": total,
        "maxPossible": max_possible,
        "champion": champion,
        "championAlive": champ_alive,
        "tiebreaker": picks.get("tiebreaker", 0),
        "f4Picks": f4_picks,
        "f4Alive": f4_alive,
    }


def build_pick_entry(raw_team, round_key, actual_team=None, status_override=None, points_override=None):
    team = clean_team_name(raw_team)
    if status_override is not None:
        status = status_override
        points_awarded = points_override if points_override is not None else 0
    elif actual_team is None:
        status = "pending" if team else "empty"
        points_awarded = 0
    elif team:
        is_correct = names_match(team, actual_team)
        status = "correct" if is_correct else "incorrect"
        points_awarded = ROUND_POINTS[round_key] if is_correct else 0
    else:
        status = "empty"
        points_awarded = 0

    return {
        "team": team,
        "round": round_key,
        "status": status,
        "pointsAwarded": points_awarded,
        "actualTeam": clean_team_name(actual_team) if actual_team else "",
    }


def build_region_comparison(region, picks=None, actual_mode=False):
    region_data = {}
    for round_key in COMPLETED_ROUNDS:
        positions = [pos for pos in get_pick_positions(region) if pos[2] == round_key]
        entries = []
        for idx, (row, col, _) in enumerate(positions):
            actual_team = ACTUALS_BY_POS.get((row, col))
            if actual_mode:
                entries.append(build_pick_entry(actual_team, round_key, status_override="correct", points_override=ROUND_POINTS[round_key]))
                continue

            round_picks = picks.get(region, {}).get(round_key, []) if picks else []
            raw_pick = round_picks[idx] if idx < len(round_picks) else ""
            entries.append(build_pick_entry(raw_pick, round_key, actual_team=actual_team))

        region_data[round_key] = entries
    return region_data


def build_future_groups(picks=None, actual_mode=False):
    if actual_mode:
        return [
            {
                "title": "Championship Matchup",
                "entries": [build_pick_entry(team, "F4", status_override="correct", points_override=ROUND_POINTS["F4"]) for team in FINALISTS],
            },
            {
                "title": "Champion",
                "entries": [build_pick_entry("TBD", "NCG", status_override="pending", points_override=0)],
            },
        ]

    f4_picks = picks.get("F4", []) if picks else []
    champion_pick = picks.get("NCG", "") if picks else ""
    finalist_actuals = [ACTUALS_BY_POS.get((row, col)) for row, col, _ in F4_POSITIONS]
    champion_actual = ACTUALS_BY_POS.get((NCG_POSITION[0], NCG_POSITION[1]))
    return [
        {
            "title": "Final Four Picks",
            "entries": [
                build_pick_entry(
                    pick,
                    "F4",
                    actual_team=finalist_actuals[idx] if idx < len(finalist_actuals) else None,
                )
                for idx, pick in enumerate(f4_picks)
            ],
        },
        {
            "title": "Champion Pick",
            "entries": [build_pick_entry(champion_pick, "NCG", actual_team=champion_actual)],
        },
    ]


def build_comparison_payload(name, picks=None, available=True, message=""):
    if not available:
        return {
            "name": name,
            "available": False,
            "message": message,
            "regions": {},
            "futureGroups": [],
        }

    return {
        "name": name,
        "available": True,
        "message": message,
        "regions": {region: build_region_comparison(region, picks=picks, actual_mode=False) for region in REGIONS},
        "futureGroups": build_future_groups(picks=picks, actual_mode=False),
    }


def build_actual_comparison_payload():
    return {
        "name": "Actual Results",
        "available": True,
        "message": "",
        "regions": {region: build_region_comparison(region, actual_mode=True) for region in REGIONS},
        "futureGroups": build_future_groups(actual_mode=True),
    }


# ===========================================================
# SHEET NAME -> PARTICIPANT NAME MAPPING
# ===========================================================
SHEET_TO_PARTICIPANT = {
    "Aaira": "Aaira", "Aappa": "Aappa", "Andrew": "Andrew",
    "Cam": "Cam Humphreys", "Ben C.": "Ben Curran", "Bob": "Bob",
    "Brian T": "Brian Thomas", "Caroline": "Caroline Burckardt",
    "Chance": "Chance", "Chuck": "Chuck", "Danny": "Danny Cantwell",
    "Dawson": "Dawson", "Deanne": "Deanne Rehmer", "Dom": "Dom Hoffman",
    "Heidi Brautigam": "Heidi", "Jeremy": "Jeremy Maxfield",
    "Jimena Lopez": "Jimena", "Jeff Platt": "Jeff Platt",
    "Ken": "Ken Burckardt", "Kennedy": "Kennedy Burckardt",
    "Kent": "Kent Rehmer", "Kevin Rehmer": "Kevin Rehmer",
    "Kristin": "Kristin Lange", "Kylie": "Kylie",
    "Larry": "Larry Magera", "Laurie L": "Laurie L.",
    "Liam": "Liam Curran", "Liz": "Liz B", "Lizzie": "Lizzie",
    "Margo": "Margo", "Marisa Cocco": "Marisa",
    "Michael Lange": "Michael Lange", "Michele Riley": "Michele Riley",
    "Michelle Moyer": "Michelle Moyer", "Mike McI": "Mike McInerney",
    "Neel": "Neel", "Norm": "Norm", "Owen C": "Owen Curran",
    "Paul": "Paul Gurgos", "Phil": "Phil Nuccio",
    "Ravtaj": "Ravtaj Singh", "Rhoda McI": "Rhoda",
    "Richard": "Richard", "Sally": "Sally Curran",
    "Stephen Moyer": "Stephen Moyer", "Todd Brautigam": "Todd B.",
    "Tony L": "Tony L.", "Tara": "Tara", "Tyler": "Tyler",
    "Wayne Lange": "Wayne Lange", "Wayne R": "Wayne Rehmer",
    "Wyatt": "Wyatt", "Yash": "Yash",
}

NO_SHEET_PARTICIPANTS = [
    "Emelio Prieto", "Garv Patel", "John Skwirblies", "Julian",
    "Nate", "Rob", "Spencer", "Surje P", "Tej Shah", "Ben",
]

# Sheets that exist but are completely empty (picks submitted outside workbook)
EMPTY_SHEET_PARTICIPANTS = {
    "Cam": "Cam Humphreys", "Bob": "Bob", "Paul": "Paul Gurgos",
    "Kristin": "Kristin Lange", "Jeff Platt": "Jeff Platt",
    "Ken": "Ken Burckardt", "Kennedy": "Kennedy Burckardt",
    "Lizzie": "Lizzie", "Caroline": "Caroline Burckardt",
}

ALL_FALLBACK_PARTICIPANTS = NO_SHEET_PARTICIPANTS + list(EMPTY_SHEET_PARTICIPANTS.values())

# Organizer's R1/R2 scores from Scores sheet
SCORES_SHEET_DATA = {
    "Richard": (25, 28), "Liz B": (27, 24), "Deanne Rehmer": (24, 26),
    "Mike McInerney": (28, 22), "Rhoda": (26, 24), "Tej Shah": (28, 22),
    "Aaira": (25, 24), "Emelio Prieto": (25, 24), "Brian Thomas": (26, 22),
    "Kevin Rehmer": (26, 22), "Kristin Lange": (24, 24), "Paul Gurgos": (24, 24),
    "Andrew": (25, 22), "Ben Curran": (23, 24), "Cam Humphreys": (25, 22),
    "Heidi": (25, 22), "Stephen Moyer": (25, 22), "Todd B.": (27, 20),
    "Dawson": (24, 22), "Phil Nuccio": (24, 22), "Surje P": (26, 20),
    "Tara": (24, 22), "Tony L.": (26, 20), "Wyatt": (26, 20),
    "Yash": (24, 22), "Ben": (23, 22), "Danny Cantwell": (25, 20),
    "Garv Patel": (23, 22), "Jeff Platt": (25, 20), "Jimena": (25, 20),
    "John Skwirblies": (25, 20), "Kent Rehmer": (25, 20), "Norm": (23, 22),
    "Rob": (27, 18), "Sally Curran": (23, 22), "Wayne Rehmer": (25, 20),
    "Larry Magera": (24, 20), "Margo": (26, 18), "Michelle Moyer": (24, 20),
    "Chuck": (25, 18), "Dom Hoffman": (25, 18), "Jeremy Maxfield": (23, 20),
    "Owen Curran": (25, 18), "Tyler": (21, 22), "Aappa": (22, 20),
    "Bob": (26, 16), "Kylie": (24, 18), "Liam Curran": (22, 20),
    "Chance": (21, 20), "Ken Burckardt": (21, 20), "Kennedy Burckardt": (23, 18),
    "Michael Lange": (23, 18), "Michele Riley": (23, 18), "Spencer": (23, 18),
    "Laurie L.": (22, 18), "Wayne Lange": (22, 18), "Julian": (23, 16),
    "Lizzie": (25, 14), "Marisa": (23, 16), "Nate": (21, 18),
    "Ravtaj Singh": (25, 14), "Caroline Burckardt": (21, 14), "Neel": (21, 10),
}


# ===========================================================
# MAIN
# ===========================================================
def main():
    print("Loading master workbook...")
    wb = openpyxl.load_workbook(MASTER_FILE, data_only=True)

    participants = []
    comparison_entries = {}
    mismatches = []

    # 1. Score participants with individual sheets
    for sheet_name, participant_name in SHEET_TO_PARTICIPANT.items():
        if sheet_name not in wb.sheetnames:
            print(f"  WARNING: Sheet '{sheet_name}' not found!")
            continue

        # Skip empty sheets - they'll be handled in fallback
        if sheet_name in EMPTY_SHEET_PARTICIPANTS:
            continue

        ws = wb[sheet_name]
        picks = extract_picks(ws)
        result = score_picks(picks)
        result["name"] = participant_name
        result["hasSheet"] = True
        comparison_entries[participant_name] = build_comparison_payload(participant_name, picks=picks, available=True)

        # Verify R1/R2 against organizer's Scores sheet
        org = SCORES_SHEET_DATA.get(participant_name)
        if org:
            org_r1_score, org_r2_score = org
            org_r1_correct = org_r1_score  # 1pt each
            org_r2_correct = org_r2_score // 2  # 2pt each
            my_r1 = result["roundCorrect"][0]
            my_r2 = result["roundCorrect"][1]
            if my_r1 != org_r1_correct or my_r2 != org_r2_correct:
                mismatches.append(
                    f"  {participant_name}: My R64={my_r1} vs Org={org_r1_correct}, "
                    f"My R32={my_r2} vs Org={org_r2_correct}"
                )

        participants.append(result)

    # 2. Participants without sheets or with empty sheets
    for name in ALL_FALLBACK_PARTICIPANTS:
        org = SCORES_SHEET_DATA.get(name)
        if not org:
            print(f"  WARNING: No score data for '{name}'")
            continue
        org_r1_score, org_r2_score = org
        r1_correct = org_r1_score
        r2_correct = org_r2_score // 2
        total = r1_correct * 1 + r2_correct * 2
        result = {
            "name": name,
            "hasSheet": False,
            "roundCorrect": [r1_correct, r2_correct, 0, 0, 0, 0],
            "roundScores": [r1_correct, r2_correct * 2, 0, 0, 0, 0],
            "regionScores": {"East": 0, "West": 0, "South": 0, "Midwest": 0},
            "totalPoints": total,
            "maxPossible": total,
            "champion": "Unknown",
            "championAlive": False,
            "tiebreaker": 0,
            "f4Picks": ["", ""],
            "f4Alive": [],
        }
        comparison_entries[name] = build_comparison_payload(
            name,
            available=False,
            message="No bracket sheet was available for this participant. Only Round of 64 and Round of 32 totals were entered in the master workbook.",
        )
        participants.append(result)

    # Sort: total pts desc, tiebreaker asc
    participants.sort(key=lambda x: (-x["totalPoints"], x.get("tiebreaker", 0)))
    for i, p in enumerate(participants):
        p["rank"] = i + 1

    # ===========================================================
    # PRIZES
    # ===========================================================
    prizes = {
        "buyIn": 10,
        "totalPot": 630,
        "placement": [
            {"place": "1st Place", "amount": 0},
            {"place": "2nd Place", "amount": 100},
            {"place": "3rd Place", "amount": 70},
            {"place": "4th Place", "amount": 40},
            {"place": "5th Place", "amount": 35},
            {"place": "6th Place", "amount": 30},
            {"place": "7th Place", "amount": 25},
        ],
        "roundBonuses": [],
        "regionBonuses": [],
        "roundBoobyPrizes": [],
        "consolationPrizes": [
            {"category": "First Final Four Team Lost", "amount": 5, "winner": "TBD"},
            {"category": "First Champion Lost", "amount": 5, "winner": "TBD"},
        ],
    }

    round_bonus_map = [("R64", "1st Round", 30), ("R32", "2nd Round", 25),
                       ("S16", "3rd Round", 25), ("E8", "4th Round", 25)]
    rd_idx_map = {"R64": 0, "R32": 1, "S16": 2, "E8": 3}

    for rd, label, amount in round_bonus_map:
        idx = rd_idx_map[rd]
        best = max(p["roundCorrect"][idx] for p in participants)
        winners = [p["name"] for p in participants if p["roundCorrect"][idx] == best]
        prizes["roundBonuses"].append({
            "round": label, "amount": amount, "winners": winners, "value": best,
            "splitAmount": round(amount / len(winners), 2) if len(winners) > 1 else amount,
        })

    booby_map = [("R64", "1st Round", 2), ("R32", "2nd Round", 2), ("S16", "3rd Round", 2)]
    for rd, label, amount in booby_map:
        idx = rd_idx_map[rd]
        worst = min(p["roundCorrect"][idx] for p in participants)
        losers = [p["name"] for p in participants if p["roundCorrect"][idx] == worst]
        prizes["roundBoobyPrizes"].append({
            "round": label, "amount": amount, "winners": losers, "value": worst,
        })

    for region in ["East", "West", "South", "Midwest"]:
        best = max(p["regionScores"].get(region, 0) for p in participants)
        winners = [p["name"] for p in participants if p["regionScores"].get(region, 0) == best]
        prizes["regionBonuses"].append({
            "region": region, "amount": 25, "winners": winners, "value": best,
            "splitAmount": round(25 / len(winners), 2) if len(winners) > 1 else 25,
        })

    fixed_prizes = 100 + 70 + 40 + 35 + 30 + 25
    round_b = sum(b["amount"] for b in prizes["roundBonuses"])
    region_b = sum(b["amount"] for b in prizes["regionBonuses"])
    booby_b = sum(b["amount"] for b in prizes["roundBoobyPrizes"])
    consol_b = 10
    first_place = 630 - fixed_prizes - round_b - region_b - booby_b - consol_b
    prizes["placement"][0]["amount"] = first_place

    # ===========================================================
    # BUILD JSON
    # ===========================================================
    games_played = len(ACTUALS_BY_POS)

    leaderboard = {
        "tournament": {
            "name": "2026 Rehmer-Bauman NCAA Tournament",
            "year": 2026,
            "lastUpdated": "2026-04-05T18:00:00Z",
            "totalGames": 63,
            "gamesPlayed": games_played,
            "roundsCompleted": 5,
            "status": "Championship matchup is set"
        },
        "scoring": {
            "round1": 1, "round2": 2, "round3": 4, "round4": 8,
            "round5": 10, "round6": 12,
            "maxScore": MAX_SCORE,
        },
        "rounds": [
            {"id": 1, "name": "Round of 64", "shortName": "R64", "games": 32, "pointsPer": 1, "completed": True},
            {"id": 2, "name": "Round of 32", "shortName": "R32", "games": 16, "pointsPer": 2, "completed": True},
            {"id": 3, "name": "Sweet 16", "shortName": "S16", "games": 8, "pointsPer": 4, "completed": True},
            {"id": 4, "name": "Elite 8", "shortName": "E8", "games": 4, "pointsPer": 8, "completed": True},
            {"id": 5, "name": "Final Four", "shortName": "F4", "games": 2, "pointsPer": 10, "completed": True},
            {"id": 6, "name": "Championship", "shortName": "NCG", "games": 1, "pointsPer": 12, "completed": False},
        ],
        "finalFour": {
            "teams": FINAL_FOUR_TEAMS,
            "title": "Championship Set",
            "semifinal1": {"team1": "UConn (East, #2)", "team2": "Illinois (South, #3)", "score1": 71, "score2": 62, "time": "Final"},
            "semifinal2": {"team1": "Michigan (Midwest, #1)", "team2": "Arizona (West, #1)", "score1": 91, "score2": 73, "time": "Final"},
            "championship": {"date": "April 6, 2026", "time": "8:50 PM ET", "team1": "UConn", "team2": "Michigan"},
            "venue": "Lucas Oil Stadium, Indianapolis",
        },
        "prizes": prizes,
        "participants": [],
    }

    comparison_data = {
        "regionOrder": REGIONS,
        "roundOrder": [
            {"key": "R64", "label": "R64", "points": ROUND_POINTS["R64"]},
            {"key": "R32", "label": "R32", "points": ROUND_POINTS["R32"]},
            {"key": "S16", "label": "S16", "points": ROUND_POINTS["S16"]},
            {"key": "E8", "label": "E8", "points": ROUND_POINTS["E8"]},
        ],
        "actual": build_actual_comparison_payload(),
        "participants": [],
    }

    for p in participants:
        parts = p["name"].split()
        avatar = (parts[0][0] + parts[-1][0]).upper() if len(parts) >= 2 else p["name"][:2].upper()
        leaderboard["participants"].append({
            "rank": p["rank"],
            "name": p["name"],
            "avatar": avatar,
            "totalPoints": p["totalPoints"],
            "maxPossible": p["maxPossible"],
            "roundScores": p["roundScores"],
            "roundCorrect": p["roundCorrect"],
            "regionScores": p["regionScores"],
            "champion": p.get("champion", "Unknown"),
            "championAlive": p.get("championAlive", False),
            "tiebreaker": p.get("tiebreaker", 0),
            "hasSheet": p.get("hasSheet", True),
        })

        comparison_entry = comparison_entries.get(p["name"], build_comparison_payload(p["name"], available=False, message="Bracket comparison unavailable."))
        comparison_data["participants"].append({
            "name": p["name"],
            "rank": p["rank"],
            "totalPoints": p["totalPoints"],
            "hasSheet": p.get("hasSheet", True),
            "available": comparison_entry["available"],
            "message": comparison_entry["message"],
            "regions": comparison_entry["regions"],
            "futureGroups": comparison_entry["futureGroups"],
        })

    with open(LEADERBOARD_FILE, "w") as f:
        json.dump(leaderboard, f, indent=2)

    with open(COMPARISON_FILE, "w") as f:
        json.dump(comparison_data, f, indent=2)

    # ===========================================================
    # SUMMARY
    # ===========================================================
    print(f"\n{'='*70}")
    print(f"LEADERBOARD: {len(participants)} participants | {games_played} games scored")
    print(f"{'='*70}")

    if mismatches:
        print(f"\n⚠️  R1/R2 Verification Mismatches ({len(mismatches)}):")
        for m in mismatches:
            print(m)

    print(f"\nTop 15:")
    for p in participants[:15]:
        alive = "✓" if p.get("championAlive") else "✗"
        champ = str(p.get("champion", "?"))[:12]
        tag = "" if p.get("hasSheet") else " [R1+R2]"
        print(f"  #{p['rank']:2d} {p['name']:<22s} {p['totalPoints']:3d} pts "
              f"(max {p['maxPossible']:3d})  {champ:<12s} [{alive}]{tag}")

    print(f"\nBottom 5:")
    for p in participants[-5:]:
        alive = "✓" if p.get("championAlive") else "✗"
        champ = str(p.get("champion", "?"))[:12]
        tag = "" if p.get("hasSheet") else " [R1+R2]"
        print(f"  #{p['rank']:2d} {p['name']:<22s} {p['totalPoints']:3d} pts "
              f"(max {p['maxPossible']:3d})  {champ:<12s} [{alive}]{tag}")

    yash = next(p for p in participants if p["name"] == "Yash")
    print(f"\nYash: #{yash['rank']} — {yash['totalPoints']} pts | "
          f"R64:{yash['roundCorrect'][0]} R32:{yash['roundCorrect'][1]} "
          f"S16:{yash['roundCorrect'][2]} E8:{yash['roundCorrect'][3]}")
    print(f"  Champ: {yash.get('champion')} (alive={yash.get('championAlive')}) | Max: {yash['maxPossible']}")
    print(f"  Regions: {yash['regionScores']}")

    print(f"\n1st Place Prize: ${first_place}")
    for b in prizes["roundBonuses"]:
        w = b["winners"][0] if len(b["winners"]) == 1 else f"{len(b['winners'])}-way tie"
        print(f"  Best {b['round']}: {w} ({b['value']} correct) — ${b['amount']}")
    for b in prizes["regionBonuses"]:
        w = b["winners"][0] if len(b["winners"]) == 1 else f"{len(b['winners'])}-way tie"
        print(f"  Best {b['region']}: {w} ({b['value']} pts) — ${b['amount']}")

    no_sheet = [p for p in participants if not p.get("hasSheet")]
    print(f"\n{len(no_sheet)} participants without pick sheets (R3/R4 not scored):")
    for p in no_sheet:
        print(f"  {p['name']}: {p['totalPoints']} pts")


if __name__ == "__main__":
    main()
